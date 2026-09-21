from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

path=Path(__file__).resolve().parents[1]/'TCC_3_BernardoBraga_final_sem_revisoes.docx'
doc=Document(path)
intro=next(p for p in doc.paragraphs if p.text.startswith('A Imagem 2 apresenta'))
intro_text='A Imagem 2 apresenta os casos de uso previstos para o EduITSM em notação UML (OMG, 2017). O ator Usuário representa o Aluno ou o Professor ao operar seu próprio ambiente. O papel adicional Professor permite acompanhar os ambientes dos alunos em modo somente leitura. As associações não têm setas e não representam ordem de execução. O ator Usuário e a fronteira são repetidos por legibilidade. A figura preserva o ambiente de demonstração do professor e não contabiliza casos de uso futuros como funcionalidades implementadas.'
nodes=intro._p.xpath('.//w:t');nodes[0].text=intro_text
for node in nodes[1:]:node.text=''
for table in doc.tables:
    for row in table.rows:
        if len(row.cells)==2 and row.cells[0].text=='Ator principal' and row.cells[1].text=='Aluno':
            ps=row.cells[1].paragraphs
            ns=ps[0]._p.xpath('.//w:t');ns[0].text='Usuário (Aluno ou Professor)'
            for node in ns[1:]:node.text=''
for number,filename in [(2,'im02_casos_de_uso_uml.png'),(10,'im10_estados_navegacao_uml.png')]:
    ps=doc.paragraphs
    i=next(i for i,p in enumerate(ps) if p.text.startswith(f'Imagem {number} - '))
    p=next(p for p in ps[i+1:i+4] if p._p.xpath('.//w:drawing'))
    rid=p._p.xpath('.//a:blip')[0].get(qn('r:embed'))
    doc.part.related_parts[rid]._blob=(path.parent/'diagramas_uml'/filename).read_bytes()
p=next(p for p in doc.paragraphs if p.text.startswith('Desenhar e alinhar o serviço:'))
full=p.text
prefix,body=full.split(':',1)
props=deepcopy(p.runs[0]._r.rPr)
# Mantém a ênfase só no rótulo, conforme os demais itens da mesma enumeração.
for r in p.runs:r.text=''
r=p.add_run(prefix+':')
if props is not None:r._r.insert(0,deepcopy(props))
r.bold=True
r=p.add_run(body)
if props is not None:r._r.insert(0,deepcopy(props))
r.bold=False

for prefix in ['Quadro 23 - Dicionário de dados','Quadro 24 - Tecnologias empregadas']:
    caption=next(p for p in doc.paragraphs if p.text.startswith(prefix))
    caption.paragraph_format.keep_with_next=True
    elem=caption._p.getnext()
    while elem is not None and elem.tag!=qn('w:tbl'):
        if elem.tag==qn('w:p'):
            pr=elem.get_or_add_pPr()
            if pr.find(qn('w:keepNext')) is None:pr.append(OxmlElement('w:keepNext'))
        elem=elem.getnext()
    assert elem is not None
    table=next(t for t in doc.tables if t._tbl is elem)
    for cell in table.rows[0].cells:
        for p in cell.paragraphs:p.paragraph_format.keep_with_next=True
    # Cabeçalho não fica isolado da primeira linha de dados.
    for cell in table.rows[1].cells:
        for p in cell.paragraphs:p.paragraph_format.keep_together=True
doc.save(path)
print('Ênfase do item e cabeçalhos dos Quadros 23 e 24 ajustados.')
