from pathlib import Path
from zipfile import ZipFile
import re
from docx import Document
from pypdf import PdfReader
from PIL import Image, ImageChops, ImageDraw, ImageFont

root = Path(r'C:\Users\berna\Documentos\TCC3')
qa = root / '.docx_qa_projeto_integrado'
doc = Document(root / 'TCC_3_BernardoBraga_final_sem_revisoes.docx')
text = '\n'.join(doc._element.xpath('//w:t/text()'))
assert not re.search('apêndice', text, flags=re.I)
assert not doc._element.xpath('//w:ins|//w:del|//w:moveFrom|//w:moveTo')
assert len(doc.inline_shapes) == 14
assert len(doc.tables) == 34
headings = [p.text for p in doc.paragraphs if p.style.name.startswith('Heading')]
assert headings.index('4.4.16 Endpoints da Interface de Programação de Aplicações') < headings.index('4.4.17 Esquema de Dados Implementado com Prisma') < headings.index('4.5 Resumo do Projeto') < headings.index('5 CONSIDERAÇÕES FINAIS')
numbers = [int(re.match(r'Quadro (\d+)', r.cells[0].text).group(1)) for r in doc.tables[0].rows]
assert numbers == list(range(1,32))
reader = PdfReader(qa / 'final.pdf')
print('PDF_PAGES', len(reader.pages))
print('PAGE_LABELS', [p.text for p in doc.paragraphs if 'Número de páginas:' in p.text or 'Number of pages:' in p.text])
for i,p in enumerate(reader.pages,1):
    txt = p.extract_text()
    if any(s in txt for s in ['4.4.16', '4.4.17','4.5 Resumo', 'REFERÊNCIAS', 'model Usuario']):
        print('LOCATION', i, txt[:130].replace('\n',' '))
    assert 'Erro! Indicador' not in txt and 'Error! Bookmark' not in txt

changed = []
for path in sorted(qa.glob('page-*.png')):
    old = root / '.docx_qa_sem_revisoes' / path.name
    if not old.exists() or ImageChops.difference(Image.open(path).convert('RGB'), Image.open(old).convert('RGB')).getbbox():
        changed.append(path)
print('CHANGED_PAGES', [p.stem for p in changed])
out = qa / 'contatos'
out.mkdir(exist_ok=True)
font = ImageFont.truetype(r'C:\Windows\Fonts\arial.ttf',18)
for start in range(0,len(changed),12):
    sheet = Image.new('RGB',(1290,2360),'#ddd')
    draw = ImageDraw.Draw(sheet)
    for i,path in enumerate(changed[start:start+12]):
        im = Image.open(path).convert('RGB'); im.thumbnail((412,556))
        x,y=(i%3)*430,(i//3)*590
        draw.text((x+8,y+4),path.stem,font=font,fill='black')
        sheet.paste(im,(x+(430-im.width)//2,y+28))
    sheet.save(out/f'contato-{start//12+1:02d}.jpg')
print('STRUCTURE_OK')
