"""Atualiza as figuras técnicas no único DOCX vigente, sem criar outra versão."""
from pathlib import Path
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
import json
import os
from zipfile import ZipFile
from PIL import Image
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches

ROOT=Path(__file__).resolve().parents[1]
PATH=ROOT/'TCC_3_BernardoBraga_final_sem_revisoes.docx'
QA=ROOT/'.docx_qa_diagramas_academicos'


def text(p,new):
    nodes=p._p.xpath('.//w:t')
    if not nodes:p.add_run(new)
    else:
        nodes[0].text=new
        for n in nodes[1:]:n.text=''


def main():
    raw=PATH.read_bytes();doc=Document(BytesIO(raw))
    QA.mkdir(exist_ok=True)
    before_text=[p.text for p in doc.paragraphs]
    revisions=doc.element.xpath('.//w:ins | .//w:del | .//w:moveFrom | .//w:moveTo')
    assert not revisions,'O documento contém revisões pendentes; preservar antes de editar.'
    changes={
      'A Imagem 2 apresenta o diagrama de casos de uso':
       'A Imagem 2 apresenta os casos de uso previstos para o EduITSM em notação UML (OMG, 2017). O ator Usuário representa o Aluno ou o Professor ao operar seu próprio ambiente. O papel adicional Professor permite acompanhar os ambientes dos alunos em modo somente leitura. As associações não têm setas e não representam ordem de execução. O ator Usuário e a fronteira são repetidos por legibilidade. A figura preserva o ambiente de demonstração do professor e não contabiliza casos de uso futuros como funcionalidades implementadas.',
      'Na ITIL V3, a estratégia era compreendida':
       'Na ITIL V3, a estratégia era compreendida como uma fase inicial e isolada do ciclo de vida do serviço. Na ITIL 4, essa lógica mudou: a estratégia permeia continuamente as atividades do Sistema de Valor do Serviço (SARAIVA; RIBEIRO, 2025). Na Imagem 3, o nó de intercalação (merge) recebe o fluxo inicial e os retornos alternativos, sem sincronizá-los. O losango inferior decide o caminho conforme as metas. As guardas [sim] e [não] são mutuamente exclusivas. O ciclo contínuo não apresenta nó final.',
      'Desenhar o serviço: criar o serviço de TI':
       'Desenhar e alinhar o serviço: criar o serviço de TI no portfólio, registrar custos, demanda e capacidade e estabelecer os vínculos com os objetivos estratégicos (UC06 a UC09).',
      'A Imagem 4 apresenta o diagrama de classes':
       'A Imagem 4 apresenta o modelo estrutural de domínio em UML, com atributos essenciais e multiplicidades. A composição indica propriedade exclusiva da parte: RegistroOperacional é parte de CenarioSimulacao e se associa a Servico sem uma segunda composição. A relação Usuario–Organizacao é uma associação simples. VinculoEstrategico é representada como classe associativa. Identificadores e detalhes de persistência foram omitidos nesta visão; estão no modelo de dados e no esquema Prisma.',
      'Para detalhar o comportamento interno da ferramenta':
       'Os diagramas de sequência das Imagens 5 e 6 especificam comportamentos planejados para as Fases 3 e 4. As linhas de vida usam a forma papel : Tipo. Chamadas síncronas têm ponta preenchida, retornos têm linha tracejada e ponta aberta, e as aut chamadas são desenhadas com retorno à própria linha de vida. Os fragmentos alt e loop representam alternativas e repetição. Rotas HTTP, controladores e instruções SQL foram omitidos para concentrar a visão nas responsabilidades de aplicação, domínio e persistência.',
      'A Imagem 5 apresenta o diagrama de sequência do UC09':
       'A Imagem 5 detalha o UC09, associado ao RF04. O serviço obtém a organização e as entidades, valida seu pertencimento ao ambiente do aluno e consulta a contribuição acumulada. No fragmento alt, cada alternativa tem seu próprio retorno: sucesso com o vínculo persistido ou erro RN06 com o saldo disponível. O projeto prevê que consulta da soma e gravação ocorram na mesma transação, com serialização por objetivo, para preservar o limite de 100% também sob concorrência. A figura pressupõe sessão válida e não detalha os fluxos de erro de autenticação e autorização.',
      'A Imagem 6 apresenta o diagrama de sequência do UC11':
       'A Imagem 6 detalha o UC11. CenarioService coordena a operação, consulta o repositório e aciona MotorSimulacao. A semente é inicializada uma única vez, antes do loop; cada iteração gera um registro e o adiciona ao lote em memória. O repositório persiste o cenário e o lote em uma transação única e devolve o resultado. A figura representa o fluxo com parâmetros e autorização válidos. O cálculo das medições é uma responsabilidade separada e não integra a geração dos dados operacionais.',
      'A Imagem 7 apresenta o modelo entidade-relacionamento':
       'A Imagem 7 apresenta o modelo lógico relacional do EduITSM em notação pé de galinha. A figura representa as treze tabelas e as quinze relações de chave estrangeira do esquema Prisma, com chaves primárias (PK), estrangeiras (FK) e restrições de unicidade (UK). As extremidades indicam as cardinalidades mínima e máxima; o vínculo opcional de indicador com objetivo_estrategico é preservado. Este modelo não é um diagrama conceitual: descreve a organização relacional dos dados. O esquema completo, com tipos e nulabilidade, está na Seção 4.4.17.',
      'A Imagem 8 apresenta um diagrama de componentes':
       'A Imagem 8 apresenta componentes, interfaces e dependências da solução. Aplicação Web e API HTTP usam os contratos compartilhados. Os serviços de domínio dependem das interfaces de repositório, realizadas pelos adaptadores Prisma. A dependência é representada por seta tracejada aberta; a realização da interface, por linha tracejada com triângulo vazado. O componente Prisma Client concentra o acesso ao banco. A figura distingue os serviços existentes na Fase 1 das responsabilidades previstas para os módulos seguintes.',
      'Enquanto a Imagem 8 descreve componentes':
       'A Imagem 9 apresenta a implantação de produção prevista. Navegador, Node.js e PostgreSQL são ambientes de execução, não artefatos. A aplicação web, o build da API e o esquema de dados são artefatos contidos nesses ambientes. Os caminhos de comunicação são linhas contínuas sem setas. HTTPS/TLS e PostgreSQL/TLS representam a configuração de produção planejada; o quadro informativo registra separadamente as portas e o uso de Docker Compose no ambiente local atual.',
      'Esta seção apresenta o protótipo de interface da ferramenta':
       'Esta seção apresenta o protótipo da interface. A Imagem 10 representa um recorte da navegação prevista como máquina de estados UML. O estado composto Sessão autenticada contém o painel e os demais estados de navegação. Eventos de seleção de tela disparam as transições; guardas restringem o acesso por perfil. A saída ou expiração da sessão retorna ao Login. Os estados compostos de estratégia, portfólio e alinhamento agrupam suas telas, com as transições internas omitidas. O percurso ilustra o uso didático, sem impor avanço automático entre telas.',
    }
    changes={k:v.replace('aut chamadas','autochamadas') for k,v in changes.items()}
    affected=[]
    for prefix,new in changes.items():
        ps=[p for p in doc.paragraphs if p.text.startswith(prefix)]
        assert len(ps)==1,(prefix,len(ps))
        text(ps[0],new);affected.append(prefix)
    # Corrige a classificação em legenda e lista de figuras sem remover campos/bookmarks.
    phrase='Modelo entidade-relacionamento conceitual em notação pé de galinha'
    replaced=0
    for p in doc.paragraphs:
        if phrase in p.text:
            nodes=p._p.xpath('.//w:t')
            whole=''.join(n.text or '' for n in nodes)
            start=whole.find(phrase);end=start+len(phrase);cursor=0
            for n in nodes:
                old=n.text or '';a,b=cursor,cursor+len(old);cursor=b
                if a<end and b>start:
                    n.text=old[:max(0,start-a)]+('Modelo lógico relacional em notação pé de galinha' if a<=start<b else '')+old[max(0,end-a):]
            replaced+=1
    assert replaced>=1
    # Preserva os atores dos casos detalhados: ambos operam o próprio ambiente.
    actor_cells=0
    files={2:'im02_casos_de_uso_uml.png',3:'im03_atividade_uml.png',4:'im04_classes_uml.png',
           5:'im05_sequencia_uc09_uml.png',6:'im06_sequencia_uc11_uml.png',7:'im07_mer_conceitual.png',
           8:'im08_componentes_uml.png',9:'im09_implantacao_uml.png',10:'im10_estados_navegacao_uml.png'}
    expected_media={}
    for number,filename in files.items():
        ps=doc.paragraphs
        indexes=[i for i,p in enumerate(ps) if p.text.startswith(f'Imagem {number} - ')]
        assert len(indexes)==1,(number,indexes)
        i=indexes[0]
        imagep=next(p for p in ps[i+1:i+4] if p._p.xpath('.//w:drawing'))
        blips=imagep._p.xpath('.//a:blip');assert len(blips)==1
        rid=blips[0].get(qn('r:embed'));part=doc.part.related_parts[rid]
        img=(ROOT/'diagramas_uml'/filename).read_bytes();part._blob=img
        expected_media[str(part.partname).lstrip('/')]=sha256(img).hexdigest()
        w,h=Image.open(BytesIO(img)).size;cx=int(Inches(6.15));cy=round(cx*h/w)
        for extent in imagep._p.xpath('.//wp:extent | .//a:xfrm/a:ext'):
            extent.set('cx',str(cx));extent.set('cy',str(cy))
        for pr in imagep._p.xpath('.//wp:docPr'):
            pr.set('name',ps[i].text);pr.set('descr',ps[i].text)
        ps[i].paragraph_format.keep_with_next=True
        imagep.paragraph_format.keep_with_next=True
        imagep.paragraph_format.line_spacing=1
        imagep.paragraph_format.space_after=0
        j=next(j for j,p in enumerate(ps) if p._p is imagep._p)
        ps[j+1].paragraph_format.keep_with_next=False
    # Referência normativa ligada ao parágrafo introdutório dos diagramas.
    ref='OBJECT MANAGEMENT GROUP (OMG). OMG Unified Modeling Language (OMG UML), version 2.5.1. 2017. Disponível em: https://www.omg.org/spec/UML/2.5.1. Acesso em: 9 set. 2026.'
    if not any(p.text.startswith('OBJECT MANAGEMENT GROUP') for p in doc.paragraphs):
        anchor=next(p for p in doc.paragraphs if p.text.startswith('SARTO, Prisilla'))
        p=anchor.insert_paragraph_before(ref)
        p.style=anchor.style
        if anchor._p.pPr is not None:
            p._p.insert(0,deepcopy(anchor._p.pPr))
        if anchor.runs and anchor.runs[0]._r.rPr is not None:p.runs[0]._r.insert(0,deepcopy(anchor.runs[0]._r.rPr))
    track=doc.settings.element.find(qn('w:trackRevisions'))
    if track is not None:doc.settings.element.remove(track)
    assert len(doc.inline_shapes)==14
    assert '4.4.16 Endpoints da Interface de Programação de Aplicações' in [p.text for p in doc.paragraphs]
    assert '4.4.17 Esquema de Dados Implementado com Prisma' in [p.text for p in doc.paragraphs]
    assert not any(p.text.startswith('APÊNDICE A') or p.text.startswith('APÊNDICE B') for p in doc.paragraphs)
    out=BytesIO();doc.save(out)
    with ZipFile(BytesIO(out.getvalue())) as z:
        for member,digest in expected_media.items():assert sha256(z.read(member)).hexdigest()==digest
    assert PATH.read_bytes()==raw,'Arquivo alterado externamente; operação cancelada.'
    # Troca atômica do arquivo: o temporário não é outro documento de entrega.
    tmp=QA/'documento-em-gravacao.tmp'
    tmp.write_bytes(out.getvalue());os.replace(tmp,PATH)
    report={'figures_updated':9,'inline_shapes':len(doc.inline_shapes),'tables':len(doc.tables),
      'actor_cells_updated':actor_cells,'paragraphs_updated':affected,'media_sha256':expected_media,
      'document_sha256':sha256(out.getvalue()).hexdigest()}
    (QA/'verificacao_estrutural.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(report,ensure_ascii=False,indent=2))


if __name__=='__main__':main()
