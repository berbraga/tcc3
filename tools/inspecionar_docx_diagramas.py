from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(r"C:\Users\berna\Documentos\TCC3")
DOCX = ROOT / "TCC_3_BernardoBraga_atualizado.docx"


doc = Document(DOCX)
for index in list(range(70, min(130, len(doc.paragraphs)))) + list(range(390, min(490, len(doc.paragraphs)))):
    text = doc.paragraphs[index].text.strip()
    if text:
        print(f"TXT P{index}: {text}")

print("--- IMAGENS ---")
for index, paragraph in enumerate(doc.paragraphs):
    drawings = paragraph._p.xpath(".//w:drawing")
    if not drawings:
        continue
    rel_ids = []
    for blip in paragraph._p.xpath(".//a:blip"):
        rel_id = blip.get(qn("r:embed"))
        if rel_id:
            rel_ids.append(rel_id)
    before = doc.paragraphs[index - 1].text if index else ""
    after = doc.paragraphs[index + 1].text if index + 1 < len(doc.paragraphs) else ""
    print(f"P{index} rel={rel_ids}")
    print(f"  before={before[:180]!r}")
    print(f"  text={paragraph.text[:180]!r}")
    print(f"  after={after[:180]!r}")
