from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import unicodedata

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches


ROOT = Path(r"C:\Users\berna\Documentos\TCC3")
SOURCE = ROOT / "TCC_3_BernardoBraga_atualizado.docx"
OUTPUT = ROOT / "TCC_3_BernardoBraga_UML_corrigido.docx"
DIAGRAMS = ROOT / "diagramas_uml"


def norm(value: str) -> str:
    value = unicodedata.normalize("NFKD", value)
    value = "".join(c for c in value if not unicodedata.combining(c))
    return " ".join(value.replace("—", "-").replace("–", "-").split()).lower()


def find_one(doc: Document, startswith: str):
    key = norm(startswith)
    matches = [p for p in doc.paragraphs if norm(p.text).startswith(key)]
    if len(matches) != 1:
        raise RuntimeError(f"Esperado um parágrafo iniciado por {startswith!r}; encontrados {len(matches)}")
    return matches[0]


def replace_paragraph(paragraph, text: str) -> None:
    rpr = None
    for run in paragraph.runs:
        if run._r.rPr is not None:
            rpr = deepcopy(run._r.rPr)
            break
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    run = paragraph.add_run(text)
    if rpr is not None:
        run._r.insert(0, rpr)


def replace_figure(doc: Document, caption_prefix: str, image_name: str, alt_text: str, width=6.15):
    caption = find_one(doc, caption_prefix)
    paragraphs = doc.paragraphs
    index = next(i for i, p in enumerate(paragraphs) if p._p is caption._p)
    image_paragraph = None
    for p in paragraphs[index + 1:index + 4]:
        if p._p.xpath(".//w:drawing"):
            image_paragraph = p
            break
    if image_paragraph is None:
        raise RuntimeError(f"Imagem após {caption_prefix!r} não encontrada")
    image_index = next(i for i, p in enumerate(paragraphs) if p._p is image_paragraph._p)
    for p in paragraphs[index:image_index]:
        p.paragraph_format.keep_with_next = True
    for child in list(image_paragraph._p):
        if child.tag != qn("w:pPr"):
            image_paragraph._p.remove(child)
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.space_before = 0
    image_paragraph.paragraph_format.space_after = 0
    image_paragraph.paragraph_format.keep_with_next = True
    run = image_paragraph.add_run()
    run.add_picture(str(DIAGRAMS / image_name), width=Inches(width))
    for doc_pr in image_paragraph._p.xpath(".//wp:docPr"):
        doc_pr.set("name", alt_text)
        doc_pr.set("descr", alt_text)


def main():
    doc = Document(SOURCE)

    replacements = {
        "Número de páginas: 121": "Número de páginas: 122",
        "Number of pages: 121": "Number of pages: 122",
        "A Imagem 2 apresenta o diagrama de casos de uso do EduITSM": (
            "A Imagem 2 apresenta o diagrama de casos de uso do EduITSM segundo a notação UML 2.x. "
            "A fronteira retangular delimita o sistema, as elipses representam os casos de uso e as linhas sem "
            "seta representam associações de participação. Como os dois perfis compartilham a maior parte das "
            "funcionalidades, o diagrama utiliza generalização de atores: Aluno e Professor especializam o ator "
            "Usuário. O ator Usuário é repetido nas duas margens apenas para reduzir cruzamentos; trata-se do mesmo "
            "elemento do modelo. O único caso de uso exclusivo é o UC14, atribuído ao Professor."
        ),
        "Na ITIL V3, a estratégia era compreendida como uma fase inicial": (
            "Na ITIL V3, a estratégia era compreendida como uma fase inicial e isolada do ciclo de vida do serviço. "
            "Na ITIL 4, essa lógica mudou: a estratégia permeia continuamente as atividades do Sistema de Valor do "
            "Serviço (SARAIVA; RIBEIRO, 2025). A Imagem 3 utiliza a notação de diagrama de atividade UML: o círculo "
            "preenchido indica o nó inicial, os retângulos arredondados representam ações, o losango representa a "
            "decisão e as expressões entre colchetes são condições de guarda. Como o processo é contínuo, o modelo "
            "retorna à análise do ambiente e não apresenta nó final."
        ),
        "A Imagem 4 apresenta o diagrama de classes do EduITSM": (
            "A Imagem 4 apresenta o diagrama de classes de domínio do EduITSM segundo a UML 2.x. Cada classe é "
            "representada por um retângulo compartimentado com nome e atributos essenciais. As associações possuem "
            "multiplicidades explícitas; os losangos preenchidos representam composições, nas quais a parte não possui "
            "ciclo de vida independente do todo. Organizacao atua como raiz de agregação do ambiente de trabalho. "
            "As operações não são repetidas neste modelo estrutural porque o comportamento dos casos críticos é "
            "detalhado nos diagramas de sequência da Seção 4.4.9."
        ),
        "Para detalhar o comportamento interno da ferramenta": (
            "Para detalhar o comportamento interno da ferramenta, foram elaborados diagramas de sequência UML para "
            "os casos de uso UC09 e UC11. As linhas de vida são classificadas pelos estereótipos actor, boundary, "
            "control, service, repository e database. Mensagens síncronas usam linha contínua, retornos usam linha "
            "tracejada, barras estreitas indicam ativação e os quadros alt e loop representam fragmentos combinados "
            "com condições de execução."
        ),
        "A Imagem 5 apresenta o diagrama de sequência do caso de uso UC09": (
            "A Imagem 5 apresenta o diagrama de sequência do UC09, responsável pelo RF04. Após o aluno submeter o "
            "serviço, o objetivo, a justificativa e a contribuição, o controlador encaminha a requisição ao serviço "
            "de domínio. Esse serviço deriva a organização do usuário autenticado, confirma que os elementos pertencem "
            "ao mesmo ambiente e consulta a contribuição já acumulada. O fragmento alt representa a RN06: o vínculo é "
            "persistido quando o total permanece em até 100%; caso contrário, a API devolve erro de regra de negócio "
            "com a contribuição ainda disponível."
        ),
        "A Imagem 6 apresenta o diagrama de sequência do caso de uso UC11": (
            "A Imagem 6 apresenta o diagrama de sequência do UC11, responsável por gerar registros operacionais "
            "simulados. O serviço valida a organização e os serviços autorizados, inicializa o gerador pseudoaleatório "
            "com a semente informada e executa o fragmento loop uma vez para cada registro solicitado. O cenário e o "
            "lote são persistidos na mesma transação. O retorno encerra o escopo do UC11 na geração dos dados brutos; "
            "a apuração das medições permanece responsabilidade separada do módulo de indicadores, preservando a "
            "distinção entre dado operacional e resultado calculado."
        ),
        "A Imagem 7 apresenta o modelo de entidade e relacionamento do banco de dados": (
            "A Imagem 7 apresenta o modelo entidade-relacionamento conceitual do EduITSM em notação pé de galinha. "
            "As entidades exibem as chaves primárias, estrangeiras e restrições de unicidade mais relevantes, enquanto "
            "os símbolos nas extremidades das relações registram participação mínima e máxima. O modelo relacional foi "
            "derivado do diagrama de classes, mas não deve ser confundido com ele: a UML descreve conceitos do domínio, "
            "enquanto o modelo de dados descreve como esses conceitos são persistidos. Na implementação, as treze "
            "entidades utilizam UUID e enumerações nativas do PostgreSQL por meio do Prisma."
        ),
        "A classe associativa VinculoEstrategico foi convertida na tabela vinculo_estrategico": (
            "A associação muitos-para-muitos entre servico e objetivo_estrategico é resolvida pela entidade associativa "
            "vinculo_estrategico, que armazena a justificativa de valor e o percentual de contribuição. A cardinalidade "
            "entre usuario e organizacao é de um para zero ou um no esquema, com restrição única sobre usuario_id; após "
            "o cadastro atômico, cada aluno possui exatamente uma organização. Relações dependentes usam exclusão em "
            "cascata, enquanto a referência opcional de indicador para objetivo usa SetNull. O Quadro 23 detalha as "
            "tabelas e o Apêndice B reproduz o esquema Prisma implementado."
        ),
        "A Imagem 8 apresenta a arquitetura lógica da solução": (
            "A Imagem 8 apresenta um diagrama de componentes UML da solução. A Aplicação Web depende da API REST por "
            "HTTP e JSON; a API delega as regras aos serviços de domínio; esses serviços dependem de interfaces de "
            "repositório; e a implementação dos repositórios utiliza o Prisma para acessar o PostgreSQL. O componente "
            "Contratos concentra os tipos e schemas Zod compartilhados. No incremento atual, autenticação e organização "
            "estão implementadas; estratégia, portfólio, simulação e indicadores permanecem planejados nas fases seguintes."
        ),
        "A separação prevista entre o motor de simulação e o módulo de cálculo": (
            "A separação prevista entre o motor de simulação e o módulo de cálculo de indicadores é uma decisão de "
            "responsabilidade arquitetural: o primeiro produzirá registros operacionais fictícios e o segundo apurará "
            "medições sobre esses registros. Essa divisão permite testar o cálculo independentemente da geração, atende "
            "ao RNF08 e mantém no software a distinção conceitual entre dado de operação e indicador estratégico."
        ),
        "Enquanto a Imagem 8 descreve a organização lógica da solução": (
            "Enquanto a Imagem 8 descreve componentes e dependências lógicas, a Imagem 9 utiliza a notação de diagrama "
            "de implantação UML para representar nós, ambientes de execução, artefatos e caminhos de comunicação. Na "
            "implantação de produção prevista, a aplicação web é executada no navegador; a API e o motor de simulação "
            "residem no servidor de aplicação; e o esquema é mantido no servidor PostgreSQL. A comunicação externa usa "
            "HTTPS/TLS na porta 443 e a conexão com o banco usa PostgreSQL/TLS na porta 5432. O bloco inferior registra "
            "a configuração atual de desenvolvimento local com Docker Compose."
        ),
        "Esta seção apresenta o protótipo de interface da ferramenta": (
            "Esta seção apresenta o protótipo de interface da ferramenta. A Imagem 10 modela a navegação como uma "
            "máquina de estados UML: cada tela é um estado, as setas representam transições e os rótulos entre colchetes "
            "representam condições de guarda. O retorno do painel de indicadores à análise de ambiente explicita o ciclo "
            "de revisão estratégica, e o acesso ao acompanhamento de alunos é condicionado ao perfil Professor."
        ),
    }
    for prefix, text in replacements.items():
        replace_paragraph(find_one(doc, prefix), text)

    captions = {
        "Imagem 2 - Diagrama de casos de uso da ferramenta EduITSM": "Imagem 2 - Diagrama de casos de uso UML da ferramenta EduITSM",
        "Imagem 3 - Diagrama de atividade do ciclo contínuo da estratégia de serviço": "Imagem 3 - Diagrama de atividade UML do ciclo contínuo da estratégia de serviço",
        "Imagem 4 - Diagrama de classes da ferramenta EduITSM": "Imagem 4 - Diagrama de classes UML do domínio EduITSM",
        "Imagem 5 - Diagrama de sequência do caso de uso UC09": "Imagem 5 - Diagrama de sequência UML do UC09 - Vincular serviço a objetivo estratégico",
        "Imagem 6 - Diagrama de sequência do caso de uso UC11": "Imagem 6 - Diagrama de sequência UML do UC11 - Gerar dados operacionais simulados",
        "Imagem 7 - Modelo de entidade e relacionamento do banco de dados": "Imagem 7 - Modelo entidade-relacionamento conceitual em notação pé de galinha",
        "Imagem 8 - Arquitetura da solução": "Imagem 8 - Diagrama de componentes UML da solução",
        "Imagem 9 - Diagrama de implantação da solução": "Imagem 9 - Diagrama de implantação UML da solução",
        "Imagem 10 - Mapa de navegação das telas": "Imagem 10 - Diagrama de máquina de estados UML da navegação",
    }
    for prefix, text in captions.items():
        replace_paragraph(find_one(doc, prefix), text)

    figures = [
        ("Imagem 2 - Diagrama de casos de uso UML", "im02_casos_de_uso_uml.png", "Diagrama de casos de uso UML do EduITSM"),
        ("Imagem 3 - Diagrama de atividade UML", "im03_atividade_uml.png", "Diagrama de atividade UML do ciclo contínuo da estratégia"),
        ("Imagem 4 - Diagrama de classes UML", "im04_classes_uml.png", "Diagrama de classes UML do domínio EduITSM"),
        ("Imagem 5 - Diagrama de sequência UML do UC09", "im05_sequencia_uc09_uml.png", "Diagrama de sequência UML do caso de uso UC09"),
        ("Imagem 6 - Diagrama de sequência UML do UC11", "im06_sequencia_uc11_uml.png", "Diagrama de sequência UML do caso de uso UC11"),
        ("Imagem 7 - Modelo entidade-relacionamento conceitual", "im07_mer_conceitual.png", "Modelo entidade-relacionamento conceitual do EduITSM"),
        ("Imagem 8 - Diagrama de componentes UML", "im08_componentes_uml.png", "Diagrama de componentes UML da solução EduITSM"),
        ("Imagem 9 - Diagrama de implantação UML", "im09_implantacao_uml.png", "Diagrama de implantação UML da solução EduITSM"),
        ("Imagem 10 - Diagrama de máquina de estados UML", "im10_estados_navegacao_uml.png", "Diagrama de máquina de estados UML da navegação do EduITSM"),
    ]
    for prefix, image_name, alt_text in figures:
        replace_figure(doc, prefix, image_name, alt_text)

    doc.core_properties.subject = "Versão atualizada do TCC3 com diagramas UML e modelo entidade-relacionamento revisados"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
