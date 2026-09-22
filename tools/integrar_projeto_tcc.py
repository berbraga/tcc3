from copy import deepcopy
from pathlib import Path
import re
from lxml import etree

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

PATH = Path(r'C:\Users\berna\Documentos\TCC3\TCC_3_BernardoBraga_final_sem_revisoes.docx')
doc = Document(PATH)
body = doc._element.body

def find(prefix):
    found = [p for p in doc.paragraphs if p.text.startswith(prefix)]
    assert len(found) == 1, (prefix, len(found))
    return found[0]

def replace_text(p, old, new):
    # Replace across split runs without disturbing bookmarks, fields or run styles.
    nodes = p.xpath('.//w:t')
    text = ''.join(n.text or '' for n in nodes)
    starts = []
    cursor = 0
    for n in nodes:
        starts.append(cursor)
        cursor += len(n.text or '')
    for match in reversed(list(re.finditer(re.escape(old), text))):
        start, end = match.span()
        for n, offset in zip(nodes, starts):
            value = n.text or ''
            limit = offset + len(value)
            if offset < end and limit > start:
                a, b = max(0, start-offset), min(len(value), end-offset)
                n.text = value[:a] + (new if offset <= start < limit else '') + value[b:]
                n.set(qn('xml:space'), 'preserve')

a = find('APÊNDICE A')
b = find('APÊNDICE B')
anchor = find('4.5 Resumo do Projeto')
reference_heading = find('4.4.15 Estado Atual da Implementação')
blocks = list(body)[body.index(a._p):-1]
assert body[-1].tag == qn('w:sectPr')
assert not any(e.xpath('.//w:sectPr') for e in blocks)
original_tables = [deepcopy(e) for e in blocks if e.tag == qn('w:tbl')]
code_start = find('generator client {')
schema_nodes = list(body)[body.index(code_start._p):-1]
schema_xml = [e.xml for e in schema_nodes]
for block in blocks:
    anchor._p.addprevious(block)

for p, text in [
    (a, '4.4.16 Endpoints da Interface de Programação de Aplicações'),
    (b, '4.4.17 Esquema de Dados Implementado com Prisma'),
]:
    replace_text(p._p, p.text, text)
    p.style = reference_heading.style
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = False

replacements = {
    'o Apêndice B': 'a Seção 4.4.17',
    'do Apêndice A': 'da Seção 4.4.16',
    'Este apêndice relaciona': 'Esta seção relaciona',
    'Este apêndice reproduz': 'Esta seção apresenta',
}
for p in body.xpath('.//w:p'):
    for old, new in replacements.items():
        replace_text(p, old, new)
    # Placeholder pass prevents cascading renumbering.
    for old, new in {31: 27, 27: 28, 28: 29, 29: 30, 30: 31}.items():
        replace_text(p, f'Quadro {old}', f'__QUADRO_{new}__')
    for number in range(27,32):
        replace_text(p, f'__QUADRO_{number}__', f'Quadro {number}')

# Keep the list in chapter order, with live page references to the captions.
listing = doc.tables[0]
rows = sorted(list(listing.rows), key=lambda r: int(re.match(r'Quadro (\d+)',r.cells[0].text).group(1)))
for row in rows:
    listing._tbl.append(row._tr)
ids = [int(e.get(qn('w:id'))) for e in body.xpath('.//w:bookmarkStart')]
bookmark_id = max(ids, default=0) + 1
for row in listing.rows:
    number = int(re.match(r'Quadro (\d+)', row.cells[0].text).group(1))
    candidates = [p for p in doc.paragraphs if re.match(rf'^Quadro {number}\s*[-–—]', p.text)]
    assert len(candidates) == 1, (number, len(candidates))
    caption = candidates[0]
    bookmark = f'ProjetoQuadro{number}'
    start = OxmlElement('w:bookmarkStart')
    start.set(qn('w:id'),str(bookmark_id)); start.set(qn('w:name'),bookmark)
    end = OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'),str(bookmark_id))
    caption._p.insert(1 if caption._p.pPr is not None else 0,start)
    caption._p.append(end)
    bookmark_id += 1
    p = row.cells[1].paragraphs[0]
    rpr = deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    for child in list(p._p):
        if child.tag != qn('w:pPr'): p._p.remove(child)
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'),f' PAGEREF {bookmark} \\h ')
    run = OxmlElement('w:r')
    if rpr is not None: run.append(rpr)
    t = OxmlElement('w:t'); t.text='0'; run.append(t); field.append(run); p._p.append(field)

# Verify that moving the sections preserved every schema line and endpoint cell.
assert schema_xml == [e.xml for e in schema_nodes]
moved_tables = [e for e in blocks if e.tag == qn('w:tbl')]
assert all(etree.tostring(x, method='c14n', exclusive=True) == etree.tostring(y, method='c14n', exclusive=True) for x,y in zip(original_tables,moved_tables))
assert body.index(a._p) < body.index(b._p) < body.index(anchor._p)
assert len(doc.inline_shapes) == 14
doc.save(PATH)
print('Atualizado no mesmo arquivo:', PATH)
print('Blocos movidos:', len(blocks))
print('Quadros renumerados e lista vinculada às páginas das legendas.')
