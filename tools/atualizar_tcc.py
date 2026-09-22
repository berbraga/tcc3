from __future__ import annotations

import sys
from copy import deepcopy
from pathlib import Path
import re
import unicodedata

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(r"C:\Users\berna\Documentos\TCC3")
SOURCE = ROOT / "TCC_2_BernardoBraga_banca.docx"
OUTPUT = ROOT / "TCC_3_BernardoBraga_atualizado.docx"
SCHEMA = ROOT / "apps" / "api" / "prisma" / "schema.prisma"


def norm(text: str) -> str:
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.replace("—", "-").replace("–", "-")
    return " ".join(text.split())


ACCENTS = {
    "Apendice": "Apêndice", "APENDICE": "APÊNDICE", "Sao": "São", "Jose": "José",
    "Servicos": "Serviços", "servicos": "serviços", "Servico": "Serviço", "servico": "serviço",
    "gestao": "gestão", "Estrategia": "Estratégia", "estrategia": "estratégia", "estrategias": "estratégias", "estrategico": "estratégico",
    "estrategica": "estratégica", "estrategicos": "estratégicos", "pratica": "prática", "praticas": "práticas",
    "decisoes": "decisões", "tecnologia": "tecnologia", "negocio": "negócio", "negocios": "negócios",
    "analise": "análise", "analises": "análises", "curriculos": "currículos", "instituicoes": "instituições",
    "educacional": "educacional", "pedagogico": "pedagógico", "pedagogica": "pedagógica", "pedagogicos": "pedagógicos",
    "avaliacao": "avaliação", "avaliacoes": "avaliações", "aplicacao": "aplicação", "aplicacoes": "aplicações",
    "graduacao": "graduação", "computacao": "computação", "area": "área", "areas": "áreas",
    "publico": "público", "publicos": "públicos", "usuarios": "usuários", "usuario": "usuário",
    "propria": "própria", "proprio": "próprio", "proprios": "próprios", "unico": "único", "unica": "única",
    "didatica": "didática", "construcao": "construção", "execucao": "execução", "integracao": "integração",
    "requisito": "requisito", "requisitos": "requisitos", "funcional": "funcional", "funcionais": "funcionais",
    "descricao": "descrição", "descricoes": "descrições", "secao": "seção", "secoes": "seções",
    "solucao": "solução", "solucoes": "soluções", "implementacao": "implementação", "implementacoes": "implementações",
    "implementada": "implementada", "implementadas": "implementadas", "implementado": "implementado", "implementados": "implementados",
    "Implementacao": "Implementação", "Organizacao": "Organização", "Organizacoes": "Organizações",
    "organizacao": "organização", "organizacoes": "organizações", "fundacao": "fundação",
    "executavel": "executável", "executaveis": "executáveis", "criacao": "criação", "ficticia": "fictícia",
    "suites": "suítes", "invalidas": "inválidas", "didaticos": "didáticos", "didatico": "didático",
    "implementara": "implementará", "atualizacao": "atualização", "atualizacoes": "atualizações",
    "primarias": "primárias", "substituicao": "substituição", "dominio": "domínio", "sao": "são",
    "derivacao": "derivação", "Relacoes": "Relações", "existencia": "existência", "autonoma": "autônoma",
    "proximas": "próximas", "decisao": "decisão", "contabiliza-los": "contabilizá-los",
    "verificacao": "verificação", "Estacio": "Estácio", "tematica": "temática", "panoramica": "panorâmica",
    "espaco": "espaço", "evidencia": "evidência", "reforcou": "reforçou", "avancou": "avançou",
    "Avaliacao": "Avaliação", "sequencia": "sequência", "questionario": "questionário",
    "apendice": "apêndice", "indices": "índices", "exigirao": "exigirão", "valido": "válido",
    "operarao": "operarão", "nativa": "nativa", "fase": "fase", "contribuicao": "contribuição",
    "capitulo": "capítulo", "especificacao": "especificação", "Secao": "Seção", "proposito": "propósito",
    "necessarios": "necessários", "ultimo": "último", "separacao": "separação", "repositorios": "repositórios",
    "Numero": "Número", "paginas": "páginas", "pagina": "página", "mantem": "mantém",
    "concluida": "concluída", "concluido": "concluído", "conclusao": "conclusão", "repositorio": "repositório",
    "monorepositorio": "monorepositório", "persistencia": "persistência", "autenticacao": "autenticação",
    "programacao": "programação", "operacao": "operação", "operacoes": "operações", "evolucao": "evolução",
    "informacoes": "informações", "relacao": "relação", "relacoes": "relações", "associacao": "associação",
    "identificadores": "identificadores", "inteiros": "inteiros", "enumeracoes": "enumerações", "migracao": "migração",
    "restricao": "restrição", "restricoes": "restrições", "exclusao": "exclusão", "indicador": "indicador",
    "referencia": "referência", "referencias": "referências", "tecnologias": "tecnologias", "navegacao": "navegação",
    "comunicacao": "comunicação", "aplicacoes": "aplicações", "aplicacao": "aplicação", "sessao": "sessão",
    "navegador": "navegador", "protegidas": "protegidas", "derivam": "derivam", "identificador": "identificador",
    "atomico": "atômico", "atomica": "atômica", "invalido": "inválido", "invalidos": "inválidos",
    "inválida": "inválida", "expirado": "expirado", "integral": "integral", "cenarios": "cenários",
    "cenario": "cenário", "verificados": "verificados", "dependem": "dependem", "continuam": "continuam",
    "estado": "estado", "atual": "atual", "direcao": "direção", "portfolio": "portfólio", "custos": "custos",
    "vinculos": "vínculos", "vinculo": "vínculo", "simulacao": "simulação", "deterministica": "determinística",
    "relatorio": "relatório", "relatorios": "relatórios", "implantacao": "implantação", "robustez": "robustez",
    "compatibilidade": "compatibilidade", "contempla": "contempla", "contemplam": "contemplam", "conteudo": "conteúdo",
    "conteudos": "conteúdos", "metodologicos": "metodológicos", "metodologica": "metodológica",
    "competencia": "competência", "competencias": "competências", "pesquisa": "pesquisa",
    "adotada": "adotada", "qualitativa": "qualitativa", "exploratoria": "exploratória", "quatro": "quatro",
    "hora": "hora", "aula": "aula", "aplicada": "aplicada", "avaliada": "avaliada", "subsequentes": "subsequentes",
    "seguintes": "seguintes", "automaticos": "automáticos", "automatizados": "automatizados", "interface": "interface",
    "interfaces": "interfaces", "painel": "painel", "inicio": "início", "apos": "após", "atraves": "através",
    "tambem": "também", "ja": "já", "nao": "não", "sera": "será", "serao": "serão",
    "esta": "esta", "estao": "estão", "tres": "três", "possui": "possui", "fonte": "fonte",
    "Autoria": "Autoria", "autoria": "autoria", "prevista": "prevista", "previsto": "previsto",
    "manutencao": "manutenção", "edicao": "edição", "edicoes": "edições", "edicao": "edição",
    "operacional": "operacional", "operacionais": "operacionais", "disponibilidade": "disponibilidade",
    "critério": "critério", "criterio": "critério", "criterios": "critérios", "conexao": "conexão",
    "codigo": "código", "codigos": "códigos", "modulos": "módulos", "modulo": "módulo",
    "padrao": "padrão", "periodo": "período", "versao": "versão", "versoes": "versões",
    "lógica": "lógica", "logica": "lógica", "possivel": "possível", "acesso": "acesso",
    "sintese": "síntese", "tecnica": "técnica", "tecnico": "técnico", "tecnicos": "técnicos",
}


def accent(text: str) -> str:
    phrase_replacements = {
        " e uma competencia": " é uma competencia",
        "ITIL e o framework": "ITIL é o framework",
        " e o framework mais": " é o framework mais",
        " nao e ensinada": " não é ensinada",
        " e tratada": " é tratada",
        "respeito a gestao": "respeito à gestao",
        "cujo escopo e deliberadamente": "cujo escopo é deliberadamente",
        "seu proposito e oferecer": "seu propósito é oferecer",
        "necessarios a sua aplicacao": "necessários à sua aplicacao",
        "restrito a pratica": "restrito à pratica",
        "restrito a origem": "restrito à origem",
        "O isolamento exigido pelo RNF06 e apoiado": "O isolamento exigido pelo RNF06 é apoiado",
        "token e obrigatorio": "token é obrigatório",
        "CORS e restrito a origem": "CORS é restrito à origem",
        "CORS e restrito à origem": "CORS é restrito à origem",
        "correspondem a continuidade": "correspondem à continuidade",
        "na sequencia, as fases": "na sequência, às fases",
        "trabalho esta organizado": "trabalho está organizado",
        "A pesquisa e aplicada": "A pesquisa é aplicada",
        "a interface e uma": "a interface é uma",
        "o trabalho e organizado": "o trabalho é organizado",
        "a migracao SQL versionada no repositorio e derivada": "a migração SQL versionada no repositório é derivada",
        "A migracao SQL versionada no repositorio e derivada": "A migração SQL versionada no repositório é derivada",
        "estava concluida": "estava concluída",
    }
    for plain, accented in phrase_replacements.items():
        text = text.replace(plain, accented)
    for plain, accented in sorted(ACCENTS.items(), key=lambda item: len(item[0]), reverse=True):
        text = re.sub(rf"\b{re.escape(plain)}\b", accented, text)
    return text


def find_one(doc: Document, startswith: str):
    matches = [p for p in doc.paragraphs if norm(p.text).startswith(startswith)]
    if len(matches) != 1:
        raise RuntimeError(f"Esperado um paragrafo iniciado por {startswith!r}; encontrados {len(matches)}")
    return matches[0]


def replace_paragraph(paragraph, text: str) -> None:
    text = accent(text)
    rpr = None
    for run in paragraph.runs:
        if run._r.rPr is not None:
            rpr = deepcopy(run._r.rPr)
            break
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    run = paragraph.add_run(text)
    if rpr is not None:
        run._r.insert(0, rpr)


def insert_before(anchor, text: str, style: str = "Normal"):
    paragraph = anchor.insert_paragraph_before(accent(text), style=style)
    return paragraph


def set_cell_text(cell, text: str) -> None:
    cell.text = accent(text)
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT


def update_front_matter(doc: Document) -> None:
    replace_paragraph(find_one(doc, "Sao Jose, Junho/2026"), "Sao Jose, Setembro/2026")
    replace_paragraph(find_one(doc, "Sao Jose, 06/2026"), "Sao Jose, 09/2026")
    replace_paragraph(find_one(doc, "Sao Jose, 11/2026"), "Sao Jose, 09/2026")
    replace_paragraph(find_one(doc, "Numero de paginas: 117"), "Numero de paginas: 121")
    replace_paragraph(find_one(doc, "Number of pages: 117"), "Number of pages: 121")

    resumo = find_one(doc, "O Gerenciamento de Servicos de TI e uma competencia")
    replace_paragraph(
        resumo,
        "O Gerenciamento de Servicos de TI e uma competencia cada vez mais exigida pelo mercado de trabalho, "
        "especialmente no que diz respeito a gestao da estrategia de servico, que envolve decisoes sobre quais "
        "servicos oferecer, como alinhar a tecnologia aos objetivos do negocio e como gerar valor por meio dos "
        "servicos de TI. A ITIL (Information Technology Infrastructure Library) e o framework mais adotado "
        "mundialmente para orientar essas praticas. Entretanto, a analise das grades curriculares e ementas de "
        "cursos superiores na area de TI de sete instituicoes de ensino superior brasileiras revelou que a gestao "
        "da estrategia de servico segundo a ITIL nao e ensinada como conteudo dedicado nos cursos analisados. "
        "Diante dessa lacuna, este trabalho desenvolve uma ferramenta educacional de ITSM, acompanhada de uma "
        "unidade de aprendizagem que organiza seu uso pedagogico, para o ensino da gestao da estrategia de servico "
        "em cursos superiores na area de TI. A pesquisa e aplicada, exploratoria e qualitativa, e utiliza o modelo "
        "ADDIE para estruturar a unidade de aprendizagem. A avaliacao de Jira Service Management, ServiceNow, "
        "Freshservice e GLPI mostrou que nenhuma das quatro ferramentas atende de forma integral e nativa aos "
        "requisitos pedagogicos definidos, com destaque para o registro dos quatro Ps da estrategia. Como resultado, "
        "foi projetada uma unidade de quatro horas-aula e especificado o EduITSM. Na etapa de Desenvolvimento do "
        "TCC3, foi concluida a primeira fase executavel da ferramenta: um monorepositorio com aplicacao React e "
        "TypeScript, interface de programacao de aplicacoes Node.js e Express, persistencia PostgreSQL por Prisma, "
        "autenticacao com JWT e bcrypt, organizacao isolada por usuario, telas de login e painel inicial e testes "
        "automatizados de servico, HTTP, integracao e interface. O esquema implementado contempla as treze entidades "
        "previstas; os requisitos funcionais restantes seguem organizados nas fases subsequentes de desenvolvimento, "
        "antes da aplicacao e da avaliacao da unidade de aprendizagem."
    )

    abstract = find_one(doc, "IT service management is an increasingly demanded competence")
    replace_paragraph(
        abstract,
        "IT service management is an increasingly demanded competence in the job market, especially service "
        "strategy management, which concerns decisions about which services to offer, how to align technology with "
        "business objectives, and how to generate value through IT services. ITIL (Information Technology "
        "Infrastructure Library) is the most widely adopted framework for guiding these practices. However, an "
        "analysis of curricula and syllabi from undergraduate IT programs at seven Brazilian higher education "
        "institutions showed that ITIL service strategy management is not taught as dedicated content in the "
        "programs examined. To address this gap, this work develops an educational ITSM tool together with a learning "
        "unit that organizes its pedagogical use in undergraduate IT education. The study is applied, exploratory, "
        "and qualitative, and uses the ADDIE model to structure the learning unit. The evaluation of Jira Service "
        "Management, ServiceNow, Freshservice, and GLPI showed that none of the four tools fully and natively meets "
        "the pedagogical requirements, especially the recording of the four Ps of strategy. As a result, a four-hour "
        "learning unit was designed and EduITSM was specified. During the Development stage of TCC3, the first "
        "executable increment of the tool was completed: a monorepo with a React and TypeScript application, a "
        "Node.js and Express application programming interface, PostgreSQL persistence through Prisma, JWT and "
        "bcrypt authentication, one isolated organization per user, login and dashboard screens, and automated "
        "service, HTTP, integration, and user-interface tests. The implemented schema includes the thirteen planned "
        "entities; the remaining functional requirements are organized into subsequent development phases before "
        "the learning unit is applied and evaluated."
    )


def update_project_chapter(doc: Document) -> None:
    replace_paragraph(
        find_one(doc, "Neste capitulo sera apresentada a modelagem do projeto"),
        "Neste capitulo sao apresentados o projeto da unidade de aprendizagem, a especificacao do EduITSM e o "
        "estado atual de seu desenvolvimento no TCC3. A descricao distingue os artefatos planejados das "
        "funcionalidades que ja possuem implementacao executavel no repositorio."
    )

    replace_paragraph(
        find_one(doc, "Diante desse cenario, optou-se por desenvolver uma ferramenta educacional propria"),
        "Diante desse cenario, optou-se por desenvolver uma ferramenta educacional propria, denominada EduITSM, "
        "cujo escopo e deliberadamente restrito a pratica de Gerenciamento da Estrategia. A ferramenta nao pretende "
        "substituir uma solucao comercial de ITSM nem cobrir as demais praticas da ITIL 4: seu proposito e oferecer, "
        "em um unico fluxo de uso, exatamente os sete requisitos funcionais do Quadro 11, acrescidos dos requisitos "
        "complementares necessarios a sua aplicacao didatica. As secoes seguintes apresentam a especificacao da "
        "ferramenta e os artefatos que orientam sua construcao. A Secao 4.4.15 registra separadamente o incremento "
        "executavel ja concluido e as funcionalidades que permanecem planejadas."
    )

    replace_paragraph(
        find_one(doc, "A Imagem 7 apresenta o modelo de entidade e relacionamento do banco de dados"),
        "A Imagem 7 apresenta o modelo de entidade e relacionamento do banco de dados, derivado do diagrama de "
        "classes. Na implementacao, esse modelo foi traduzido para o esquema Prisma sobre PostgreSQL e manteve as "
        "treze entidades previstas. As chaves primarias e estrangeiras utilizam UUID, em substituicao aos identificadores "
        "inteiros do roteiro SQL preliminar do TCC2, e as enumeracoes de dominio sao materializadas como tipos nativos "
        "do PostgreSQL."
    )

    replace_paragraph(
        find_one(doc, "A classe associativa VinculoEstrategico foi convertida na tabela vinculo_estrategico"),
        "A classe associativa VinculoEstrategico foi convertida na tabela vinculo_estrategico, que resolve o "
        "relacionamento muitos-para-muitos entre servico e objetivo_estrategico e abriga os atributos proprios da "
        "associacao. O isolamento exigido pelo RNF06 e apoiado pela relacao um-para-um entre usuario e organizacao, "
        "pela restricao unica sobre organizacao.usuario_id e pela derivacao do contexto de acesso a partir do usuario "
        "autenticado. Relacoes dependentes usam exclusao em cascata quando nao possuem existencia autonoma; a "
        "referencia opcional de indicador para objetivo usa SetNull. O Quadro 23 apresenta a finalidade das tabelas, "
        "e o Apendice B reproduz o esquema Prisma implementado."
    )

    replace_paragraph(
        find_one(doc, "A Imagem 8 apresenta a arquitetura da solucao, organizada em tres camadas"),
        "A Imagem 8 apresenta a arquitetura lógica da solução, organizada em três camadas. A camada de apresentação "
        "é uma aplicação de página única executada no navegador do aluno ou do professor. A camada de aplicação é "
        "um servidor Node.js que expõe uma interface de programação de aplicações REST e foi projetado para concentrar "
        "as regras de negócio, a autenticação, o motor de simulação e o cálculo dos indicadores. A camada de dados é "
        "o banco relacional PostgreSQL. No incremento atual, estão implementados os componentes de autenticação e "
        "organização; os módulos de simulação e cálculo permanecem nas fases seguintes."
    )
    replace_paragraph(
        find_one(doc, "A separacao entre o motor de simulacao e o modulo de calculo de indicadores e intencional"),
        "A separação prevista entre o motor de simulação e o módulo de cálculo de indicadores é intencional: o "
        "primeiro produzirá os registros operacionais fictícios, e o segundo apurará os indicadores sobre esses "
        "registros. Essa divisão permitirá testar o cálculo independentemente da geração dos dados, atendendo ao "
        "RNF08, e reproduzirá na arquitetura a distinção conceitual entre dado de operação e indicador estratégico. "
        "O Quadro 24 apresenta as tecnologias adotadas e planejadas para cada finalidade."
    )
    replace_paragraph(
        find_one(doc, "Enquanto a Imagem 8 descreve a organizacao logica da solucao, a Imagem 9 apresenta o diagrama de implantacao"),
        "Enquanto a Imagem 8 descreve a organização lógica da solução, a Imagem 9 apresenta a implantação prevista. "
        "No ambiente de produção, a ferramenta será distribuída entre o navegador do aluno ou do professor, o servidor "
        "de aplicação em nuvem e o servidor PostgreSQL. A comunicação externa ocorrerá por HTTPS na porta 443, e a "
        "conexão entre aplicação e banco usará a porta 5432 com transporte cifrado. Na Fase 1, a execução e a persistência "
        "são locais, apoiadas pelo Docker Compose."
    )
    replace_paragraph(
        find_one(doc, "Essa distribuicao atende diretamente a dois requisitos nao funcionais"),
        "A distribuição proposta atende ao RNF03 e ao RNF04 porque exige do usuário apenas um navegador atual. O "
        "RNF07 será validado na Fase 6, após a hospedagem em nuvem e a verificação de disponibilidade durante a janela "
        "de aplicação da unidade de aprendizagem."
    )

    tech_table = next(t for t in doc.tables if norm(t.cell(0, 0).text) == "Camada / Finalidade")
    updates = {
        "Interface": "React 19.1.1 com TypeScript 5.9.2",
        "Navegacao": "React Router 7.18.3",
        "Comunicacao": "Axios 1.20.0 com TanStack Query 5.87.1",
        "Servidor de aplicacao": "Node.js 22 com Express 5.1.0",
        "Autenticacao": "JSON Web Token 9.0.2 e bcryptjs 3.0.2",
        "Acesso a dados": "Prisma ORM 6.12.0",
        "Banco de dados": "PostgreSQL 16",
        "Testes": "Vitest 3.2.4, Supertest 7.1.4 e Testing Library",
        "Versionamento e hospedagem": "Git, Docker Compose e hospedagem em nuvem"
    }
    for row in tech_table.rows[1:]:
        key = norm(row.cells[0].text)
        if key in updates:
            set_cell_text(row.cells[1], updates[key])
        if key == "Versionamento e hospedagem":
            set_cell_text(
                row.cells[2],
                "O versionamento Git e o ambiente local com Docker Compose ja integram o repositorio; a hospedagem "
                "em nuvem permanece prevista para a fase de robustez e implantacao, quando sera validado o RNF07."
            )

    # Find the source paragraph immediately after the technology table.
    architecture_caption = None
    body = list(doc.element.body.iterchildren())
    tech_index = body.index(tech_table._tbl)
    for child in body[tech_index + 1:]:
        if child.tag == qn("w:p"):
            from docx.text.paragraph import Paragraph
            candidate = Paragraph(child, doc)
            if norm(candidate.text).startswith("Fonte: Autoria propria"):
                architecture_caption = candidate
                break
    if architecture_caption is None:
        raise RuntimeError("Fonte do quadro de tecnologias nao encontrada")
    insert_before(
        architecture_caption,
        "A estrutura executavel foi organizada como monorepositorio npm, com os modulos apps/api, apps/web e "
        "packages/shared. Este ultimo concentra os contratos Zod e os tipos compartilhados entre cliente e servidor. "
        "A API segue a separacao entre rotas, servicos de dominio e repositorios Prisma; a aplicacao de pagina unica "
        "consome a API REST e mantem a sessao somente no sessionStorage do navegador."
    )

    test_source = None
    test_table = next(t for t in doc.tables if norm(t.cell(0, 0).text) == "Identificador")
    body = list(doc.element.body.iterchildren())
    test_index = body.index(test_table._tbl)
    for child in body[test_index + 1:]:
        if child.tag == qn("w:p"):
            from docx.text.paragraph import Paragraph
            candidate = Paragraph(child, doc)
            if norm(candidate.text).startswith("Fonte: Autoria propria"):
                test_source = candidate
                break
    if test_source is None:
        raise RuntimeError("Fonte do plano de testes nao encontrada")
    insert_before(
        test_source,
        "Na Fase 1 foram implementadas suites automatizadas com Vitest, Supertest e Testing Library. Elas cobrem "
        "servicos de autenticacao e organizacao, respostas HTTP, integracao com PostgreSQL em esquema exclusivo de "
        "teste e os estados das telas T01 e T02. Entre os cenarios verificados estao o hash de senha, o cadastro "
        "atomico do aluno e de sua organizacao, credenciais invalidas, token ausente, invalido ou expirado, bloqueio "
        "de acesso a outra organizacao, carregamento e erro no painel e edicao dos dados da organizacao. Os casos "
        "TS01 a TS08 e TS11 a TS15 dependem das funcionalidades das fases seguintes e continuam no plano de testes."
    )

    summary_heading = find_one(doc, "4.5 Resumo do Projeto")
    section = [
        ("4.4.15 Estado Atual da Implementacao", "Heading 3"),
        (
            "Em 8 de setembro de 2026, a Fase 1 do desenvolvimento do EduITSM estava concluida no repositorio. "
            "O incremento entrega uma fundacao executavel para os requisitos RF08 e RF09 e para a RN01: cadastro "
            "publico de aluno, autenticacao, criacao de uma organizacao ficticia por usuario, consulta e edicao da "
            "propria organizacao, tela T01 de login e tela T02 de painel inicial.",
            "Normal"
        ),
        (
            "O cadastro normaliza o e-mail, aplica hash bcrypt a senha e cria o usuario e a organizacao na mesma "
            "operacao de persistencia. A autenticacao emite JWT com o identificador do usuario e o perfil; as rotas "
            "protegidas derivam a organizacao desse identificador, sem aceitar um organizacaoId fornecido pelo cliente "
            "como autoridade. O segredo do token e obrigatorio, o CORS e restrito a origem configurada e a interface "
            "armazena a sessao apenas durante a aba do navegador.",
            "Normal"
        ),
        (
            "O painel inicial apresenta os dados reais da organizacao e resume as quantidades de servicos, objetivos, "
            "versao vigente da estrategia e registros operacionais. Como as rotas dos demais modulos ainda nao foram "
            "implementadas, os respectivos itens de navegacao permanecem desabilitados e identificados como proximas "
            "fases, evitando que a interface aparente funcionalidades inexistentes.",
            "Normal"
        ),
        (
            "O banco de dados ja materializa as treze entidades e suas relacoes por meio do Prisma, com migracao "
            "versionada para PostgreSQL. Essa decisao prepara os modulos subsequentes sem contabiliza-los como "
            "entregues. Os endpoints ativos da Fase 1 sao o health check, o registro, o login e a consulta e atualizacao "
            "da organizacao autenticada; os demais endpoints do Apendice A permanecem como contrato-alvo.",
            "Normal"
        ),
        (
            "As Fases 2 a 6 foram detalhadas em um roteiro de implementacao. A Fase 2 cobre a direcao estrategica e "
            "as telas T03 a T05; a Fase 3, o portfolio, custos, demanda, vinculos e indicadores; a Fase 4, a simulacao "
            "deterministica e o painel de avaliacao; a Fase 5, o relatorio e o acompanhamento pelo professor; e a Fase "
            "6, a robustez, os testes de desempenho e compatibilidade, a implantacao e a verificacao integral do fluxo.",
            "Normal"
        )
    ]
    for text, style in section:
        insert_before(summary_heading, text, style)

    replace_paragraph(
        find_one(doc, "As fases de Desenvolvimento, Implementacao e Avaliacao do modelo ADDIE serao executadas"),
        "A fase de Desenvolvimento do modelo ADDIE foi iniciada no TCC3 com a conclusao da Fase 1 do EduITSM e "
        "com o planejamento detalhado das Fases 2 a 6. O incremento atual torna executaveis a autenticacao e o "
        "ambiente isolado por usuario, enquanto preserva no banco a estrutura das treze entidades previstas. A "
        "continuidade do desenvolvimento implementara o fluxo estrategico restante e os materiais didaticos; em "
        "seguida, a unidade sera aplicada em uma turma real e seus resultados serao avaliados."
    )


def update_conclusion(doc: Document) -> None:
    replace_paragraph(
        find_one(doc, "Na analise do estado da pratica, foi realizada uma pesquisa nas grades curriculares"),
        "Na analise do estado da pratica, foi realizada uma pesquisa nas grades curriculares e ementas de sete "
        "instituicoes de ensino superior brasileiras (UNIVALI, USP, UFSC, IME, ITA, Estacio e UNISUL). Os resultados "
        "mostraram que a gestao da estrategia de servico segundo a ITIL nao e ensinada como conteudo dedicado em "
        "nenhum dos cursos investigados. Quando a tematica aparece, e tratada de forma panoramica ou compartilha "
        "espaco com outros frameworks em disciplinas mais amplas. Essa evidencia reforcou a lacuna curricular que "
        "motivou a proposta."
    )

    next_heading = find_one(doc, "5.1 PROXIMAS ETAPAS")
    insert_before(
        next_heading,
        "No TCC3, a fase de Desenvolvimento avancou com a primeira versao executavel do EduITSM. Foram implementados "
        "o monorepositorio, a API REST, a aplicacao web, o esquema relacional completo, a autenticacao, o isolamento "
        "da organizacao por usuario e as telas T01 e T02, acompanhados de testes automatizados. Esse resultado reduz "
        "o risco tecnico do projeto e estabelece a base sobre a qual serao implementados o fluxo estrategico, o motor "
        "de simulacao, os relatorios e o acompanhamento do professor."
    )

    replace_paragraph(
        find_one(doc, "As proximas etapas deste trabalho correspondem as fases de Desenvolvimento, Implementacao e Avaliacao"),
        "As proximas etapas correspondem a continuidade da fase de Desenvolvimento e, na sequencia, as fases de "
        "Implementacao e Avaliacao do ADDIE. No software, o trabalho esta organizado nas Fases 2 a 6, que completam "
        "os requisitos RF01 a RF07 e RF10 a RF13 e as regras RN02 a RN11."
    )
    replace_paragraph(
        find_one(doc, "Na fase de Desenvolvimento, prevista para julho e agosto de 2026"),
        "Na continuidade da fase de Desenvolvimento, serao implementadas a analise de ambiente, os quatro Ps e "
        "seu versionamento, os objetivos estrategicos, o portfolio, os custos, a demanda, os vinculos, os indicadores, "
        "o motor de simulacao, os relatorios e o acompanhamento em modo somente leitura pelo professor. Em paralelo, "
        "serao produzidos e revisados os slides, o texto de apoio, o estudo de caso, os roteiros de atividades, o "
        "questionario e o manual de uso, mantendo o alinhamento com os objetivos de aprendizagem definidos no Design."
    )


def update_appendices(doc: Document) -> None:
    appendix_a_intro = find_one(doc, "Este apendice relaciona os endpoints da interface de programacao de aplicacoes do EduITSM")
    replace_paragraph(
        appendix_a_intro,
        "Este apendice relaciona o contrato-alvo da interface de programacao de aplicacoes REST do EduITSM e o "
        "requisito atendido por cada rota. No incremento atual, estao implementados o health check, o registro, o "
        "login e a consulta e atualizacao da organizacao autenticada. As demais rotas continuam planejadas para as "
        "Fases 2 a 5. Todas as rotas de dominio exigirao JWT valido e operarao sobre a organizacao derivada do usuario "
        "autenticado, conforme o RNF06."
    )

    endpoint_table = next(t for t in doc.tables if norm(t.cell(0, 0).text) == "Metodo")
    implemented = {
        "/api/v1/auth/registro",
        "/api/v1/auth/login",
        "/api/v1/organizacoes/minha"
    }
    for row in endpoint_table.rows[1:]:
        route = norm(row.cells[1].text)
        description = norm(row.cells[2].text)
        suffix = " Implementado na Fase 1." if route in implemented else " Planejado para as fases seguintes."
        if "Implementado na Fase 1" not in description and "Planejado para as fases seguintes" not in description:
            set_cell_text(row.cells[2], description + suffix)

    heading = find_one(doc, "APENDICE B - SCRIPT DE CRIACAO DO BANCO DE DADOS")
    replace_paragraph(heading, "APENDICE B — ESQUEMA DE DADOS IMPLEMENTADO COM PRISMA")
    intro = find_one(doc, "Este apendice apresenta o script de criacao do banco de dados correspondente")
    replace_paragraph(
        intro,
        "Este apendice reproduz o esquema Prisma utilizado pela implementacao do EduITSM. O arquivo gera o cliente "
        "de acesso, define o PostgreSQL como banco de dados e materializa as treze entidades, os tipos enumerados, "
        "as chaves UUID, os indices e as regras referenciais. A migracao SQL versionada no repositorio e derivada "
        "deste esquema."
    )

    paragraphs = doc.paragraphs
    hidx = next(i for i, paragraph in enumerate(paragraphs) if paragraph._p is heading._p)
    source = None
    for p in paragraphs[hidx + 1:]:
        if norm(p.text).startswith("Fonte: Autoria propria"):
            source = p
            break
    if source is None:
        raise RuntimeError("Fonte do Apendice B nao encontrada")
    paragraphs = doc.paragraphs
    iidx = next(i for i, paragraph in enumerate(paragraphs) if paragraph._p is intro._p)
    sidx = next(i for i, paragraph in enumerate(paragraphs) if paragraph._p is source._p)
    for paragraph in list(paragraphs[iidx + 1:sidx]):
        paragraph._element.getparent().remove(paragraph._element)

    schema_lines = SCHEMA.read_text(encoding="utf-8").splitlines()
    for line in schema_lines:
        paragraph = source.insert_paragraph_before(line if line else " ", style="Normal")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        for run in paragraph.runs:
            run.font.name = "Courier New"
            run.font.size = Pt(8.5)
            run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Courier New")
            run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Courier New")


def main() -> None:
    if not SOURCE.exists() or not SCHEMA.exists():
        raise FileNotFoundError("Documento-fonte ou schema Prisma nao encontrado")
    doc = Document(SOURCE)
    update_front_matter(doc)
    update_project_chapter(doc)
    update_conclusion(doc)
    update_appendices(doc)
    for node in doc.element.xpath(".//w:instrText"):
        if node.text and node.text.strip().startswith("TOC "):
            node.text = ' TOC \\o "1-4" \\h \\z \\u '
    doc.core_properties.title = "Uma ferramenta educacional de ITSM para o ensino de gestao da estrategia de servicos de TI baseada na ITIL 4"
    doc.core_properties.subject = "Versao atualizada para o TCC3 conforme a Fase 1 do EduITSM"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
